"""Excel reconciliation workbook (openpyxl)."""
from __future__ import annotations

import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from reporting import metrics

HDR_FILL = PatternFill("solid", fgColor="1F4E5F")
HDR_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows
                                     if i <= len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 46)
    ws.freeze_panes = "A2"
    return ws


def export_excel(conn: sqlite3.Connection, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)

    t = metrics.campaign_totals(conn)
    camp = conn.execute("SELECT * FROM campaigns LIMIT 1").fetchone()
    summary = [("Client", camp["client_name"] if camp else ""),
               ("Campaign", camp["campaign_name"] if camp else ""),
               ("Creators", t["creators"]),
               ("Content pieces (latest snapshots)", t["content_count"]),
               ("Deliverables required / published",
                f"{t['deliverables_required']} / {t['deliverables_published']}"),
               ("Total views", t["views"]), ("Total reach", t["reach"]),
               ("Total impressions", t["impressions"]),
               ("Total engagements", t["engagements"]),
               ("Engagement rate %", t["engagement_rate_pct"]),
               ("Contracted fees", t["contracted_fees"]),
               ("Invoiced", t["invoiced"]), ("Paid", t["paid"]),
               ("Outstanding", t["outstanding"]),
               ("Open exceptions", t["open_exceptions"])]
    _sheet(wb, "Campaign Summary", ["Metric", "Value"], summary)

    names = {r["creator_id"]: r for r in conn.execute("SELECT * FROM creators")}
    creator_rows = []
    for cid, n in sorted(names.items()):
        perf = next((p for p in metrics.creator_performance(conn)
                     if p["creator_id"] == cid), {})
        fin = next((f for f in metrics.financial_table(conn)
                    if f["creator_id"] == cid), {})
        creator_rows.append([cid, n["canonical_name"], n["instagram_handle"],
                             perf.get("content_count", 0), perf.get("views", 0),
                             perf.get("reach", 0), perf.get("engagement_rate_pct"),
                             fin.get("contract_fee"), fin.get("invoice"),
                             fin.get("payout"), fin.get("variance"), fin.get("status")])
    _sheet(wb, "Creators",
           ["Creator ID", "Name", "Instagram", "Posts", "Views", "Reach",
            "Engagement Rate %", "Contract Fee", "Invoiced", "Paid", "Variance",
            "Financial Status"], creator_rows)

    _sheet(wb, "Deliverables",
           ["Creator", "Required", "Published", "Missing", "Late", "Approval Statuses",
            "Source Contract"],
           [[d["creator"], d["required"], d["published"], d["missing"], d["late"],
             d["approval_statuses"], d["source_file"]]
            for d in metrics.deliverable_table(conn)])

    _sheet(wb, "Analytics",
           ["Content ID", "Creator", "Platform", "Type", "Publish Date", "Views",
            "Reach", "Impressions", "Likes", "Comments", "Saves", "Shares",
            "Engagements", "ER %", "Captured At", "Source File", "Source Type"],
           [[c["content_id"], c["creator"], c["platform"], c["content_type"],
             c["publish_date"], c["views"], c["reach"], c["impressions"], c["likes"],
             c["comments"], c["saves"], c["shares"], c["engagements"],
             c["engagement_rate_pct"], c["captured_at"] if "captured_at" in c.keys() else "", c["source_file"], c["source_type"]]
            for c in metrics.content_performance(conn)])

    _sheet(wb, "Financial Reconciliation",
           ["Creator", "Contract Fee", "Invoice", "Payout", "Variance",
            "Outstanding", "Status"],
           [[f["creator"], f["contract_fee"], f["invoice"], f["payout"],
             f["variance"], f["outstanding"], f["status"]]
            for f in metrics.financial_table(conn)])

    exc = conn.execute("SELECT * FROM exceptions ORDER BY severity, exception_id").fetchall()
    _sheet(wb, "Exceptions",
           ["ID", "Type", "Severity", "Status", "Creator", "Description",
            "Expected", "Actual", "Evidence", "Source Files"],
           [[e["exception_id"], e["exception_type"], e["severity"], e["status"],
             names.get(e["creator_id"], {}).get("canonical_name") if e["creator_id"] else "",
             e["description"], e["expected_value"], e["actual_value"],
             e["evidence_json"], e["source_files"]] for e in exc])

    wb.save(path)
    return path
